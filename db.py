"""
Data layer for the Member Database & Dues Management System.
Pure Python + SQLite + pandas (no Streamlit imports) so it can be tested on its own.
"""
from __future__ import annotations

import io
import json
import sqlite3
import datetime as dt
from pathlib import Path

import pandas as pd

APP_DIR = Path(__file__).parent
DB_PATH = APP_DIR / "members.db"
UNI_FILE = APP_DIR / "Chongqing_Universities_Locations.xlsx"

# --------------------------------------------------------------------------------------
# Fixed choice lists used in dropdowns
# --------------------------------------------------------------------------------------
GENDERS = ["Male", "Female", "Prefer not to say"]
DEGREE_LEVELS = ["Language Program (Chinese)", "Diploma / Vocational", "Bachelor's",
                 "Master's", "PhD", "Post-Doctoral", "Exchange / Visiting", "Staff / Faculty", "Other"]
SCHOLARSHIPS = ["Chinese Government Scholarship (CSC)", "Provincial / City Scholarship",
                "University Scholarship", "Confucius Institute Scholarship",
                "Home Government Scholarship", "Self-Sponsored", "Other"]
RESIDENCE = ["On-campus dormitory", "Off-campus apartment", "Other"]
STATUSES = ["Active", "Inactive", "Graduated", "Suspended", "Pending approval"]
PENDING = "Pending approval"
EXIT_STATUSES = ["Inactive", "Graduated", "Suspended"]
PAY_METHODS = ["Cash", "WeChat Pay", "Alipay", "Bank Transfer", "Mobile Money", "Other"]

DEFAULT_SETTINGS = {
    "org_name": "Members Association – Chongqing",
    "currency": "¥",
    "membership_types": json.dumps({"Regular": 50, "Executive": 50, "Associate": 30, "Honorary": 0}),
    "due_day": "10",  # day of month dues are expected by
    "app_url": "",
}

MEMBER_FIELDS = [
    "full_name", "gender", "date_of_birth", "nationality", "phone", "email", "wechat_id",
    "university_id", "student_id", "degree_level", "field_of_study", "enrollment_year",
    "expected_graduation", "scholarship", "residence", "membership_type", "monthly_due",
    "join_date", "status", "exit_date", "emergency_name", "emergency_phone", "notes",
]

SCHEMA = """
CREATE TABLE IF NOT EXISTS universities(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name_en TEXT UNIQUE NOT NULL,
    name_cn TEXT, ownership TEXT, district TEXT, address TEXT, postal_code TEXT
);
CREATE TABLE IF NOT EXISTS members(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    member_code TEXT UNIQUE,
    full_name TEXT NOT NULL, gender TEXT, date_of_birth TEXT, nationality TEXT,
    phone TEXT, email TEXT, wechat_id TEXT,
    university_id INTEGER REFERENCES universities(id),
    student_id TEXT, degree_level TEXT, field_of_study TEXT,
    enrollment_year INTEGER, expected_graduation INTEGER,
    scholarship TEXT, residence TEXT,
    membership_type TEXT, monthly_due REAL NOT NULL DEFAULT 0,
    join_date TEXT NOT NULL, status TEXT DEFAULT 'Active', exit_date TEXT,
    emergency_name TEXT, emergency_phone TEXT, notes TEXT,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP, updated_at TEXT
);
CREATE TABLE IF NOT EXISTS payments(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    member_id INTEGER NOT NULL REFERENCES members(id) ON DELETE CASCADE,
    period TEXT NOT NULL,             -- 'YYYY-MM' the payment covers
    amount REAL NOT NULL,
    payment_date TEXT NOT NULL,
    method TEXT, reference TEXT, received_by TEXT, notes TEXT,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE INDEX IF NOT EXISTS ix_pay_member ON payments(member_id);
CREATE INDEX IF NOT EXISTS ix_pay_period ON payments(period);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
"""

MEMBER_COLS = ["id", "member_code", *MEMBER_FIELDS, "created_at", "updated_at"]
PAYMENT_COLS = ["id", "member_id", "period", "amount", "payment_date", "method",
                "reference", "received_by", "notes", "created_at"]
UNI_COLS = ["id", "name_en", "name_cn", "ownership", "district", "address", "postal_code"]


# --------------------------------------------------------------------------------------
# Connection / setup
# --------------------------------------------------------------------------------------
def connect(path: Path | str | None = None) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path or DB_PATH), check_same_thread=False)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA)
    for k, v in DEFAULT_SETTINGS.items():
        conn.execute("INSERT OR IGNORE INTO settings(key, value) VALUES (?, ?)", (k, v))
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM universities").fetchone()[0] == 0 and UNI_FILE.exists():
        import_universities(conn, UNI_FILE)
    return conn


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():  # postal codes read as floats
        s = s[:-2]
    return s or None


def import_universities(conn, file) -> int:
    """Load / refresh the university list from the Excel file. Returns rows upserted."""
    df = pd.read_excel(file)
    df.columns = [str(c).strip().lower() for c in df.columns]

    def pick(*keys):
        for c in df.columns:
            if any(k in c for k in keys):
                return c
        return None

    c_en = pick("english", "university")
    c_cn = pick("chinese")
    c_own = pick("ownership", "type")
    c_dist = pick("district", "area")
    c_addr = pick("address")
    c_post = pick("postal", "zip")
    if c_en is None:
        raise ValueError("Could not find a university name column in the file.")
    n = 0
    for _, r in df.iterrows():
        name = _clean(r[c_en])
        if not name:
            continue
        vals = (name,
                _clean(r[c_cn]) if c_cn else None,
                _clean(r[c_own]) if c_own else None,
                _clean(r[c_dist]) if c_dist else None,
                _clean(r[c_addr]) if c_addr else None,
                _clean(r[c_post]) if c_post else None)
        conn.execute(
            """INSERT INTO universities(name_en, name_cn, ownership, district, address, postal_code)
               VALUES (?,?,?,?,?,?)
               ON CONFLICT(name_en) DO UPDATE SET name_cn=excluded.name_cn, ownership=excluded.ownership,
               district=excluded.district, address=excluded.address, postal_code=excluded.postal_code""",
            vals)
        n += 1
    conn.commit()
    return n


# --------------------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------------------
def get_settings(conn) -> dict:
    s = dict(conn.execute("SELECT key, value FROM settings").fetchall())
    s["membership_types"] = json.loads(s.get("membership_types") or "{}")
    return s


def save_setting(conn, key, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value)
    conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                 (key, str(value)))
    conn.commit()


# --------------------------------------------------------------------------------------
# Universities
# --------------------------------------------------------------------------------------
def universities(conn) -> pd.DataFrame:
    return pd.read_sql("SELECT * FROM universities ORDER BY name_en", conn)


def add_university(conn, name_en, name_cn, ownership, district, address, postal_code):
    conn.execute("INSERT INTO universities(name_en,name_cn,ownership,district,address,postal_code) VALUES(?,?,?,?,?,?)",
                 (name_en.strip(), name_cn, ownership, district, address, postal_code))
    conn.commit()


# --------------------------------------------------------------------------------------
# Members
# --------------------------------------------------------------------------------------
def _next_code(conn) -> str:
    row = conn.execute("SELECT MAX(id) FROM members").fetchone()[0] or 0
    while True:
        row += 1
        code = f"MBR-{row:04d}"
        if not conn.execute("SELECT 1 FROM members WHERE member_code=?", (code,)).fetchone():
            return code


def _norm(data: dict) -> dict:
    out = {}
    for k in MEMBER_FIELDS:
        v = data.get(k)
        if isinstance(v, (dt.date, dt.datetime)):
            v = v.isoformat()[:10]
        if isinstance(v, str):
            v = v.strip() or None
        out[k] = v
    return out


def validate_member(conn, data: dict, member_id: int | None = None) -> list[str]:
    errs = []
    if not (data.get("full_name") or "").strip():
        errs.append("Full name is required.")
    if not data.get("university_id"):
        errs.append("Please select a university.")
    if not data.get("join_date"):
        errs.append("Join date is required.")
    email = (data.get("email") or "").strip()
    if email and ("@" not in email or "." not in email.split("@")[-1]):
        errs.append("Email address does not look valid.")
    if (data.get("monthly_due") or 0) < 0:
        errs.append("Monthly due cannot be negative.")
    if data.get("status") in EXIT_STATUSES and not data.get("exit_date"):
        errs.append("Please give the date the member left / graduated (billing stops then).")
    if data.get("exit_date") and data.get("join_date") and str(data["exit_date"]) < str(data["join_date"]):
        errs.append("Exit date cannot be before join date.")
    sid = (data.get("student_id") or "").strip()
    if sid:
        q = "SELECT member_code FROM members WHERE student_id=? AND university_id=?"
        args = [sid, data.get("university_id")]
        if member_id:
            q += " AND id<>?"
            args.append(member_id)
        hit = conn.execute(q, args).fetchone()
        if hit:
            errs.append(f"Student ID {sid} is already registered at this university ({hit[0]}).")
    return errs


def add_member(conn, data: dict) -> str:
    d = _norm(data)
    code = _next_code(conn)
    cols = ["member_code", *MEMBER_FIELDS]
    conn.execute(f"INSERT INTO members({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                 [code, *[d[k] for k in MEMBER_FIELDS]])
    conn.commit()
    return code


def update_member(conn, member_id: int, data: dict):
    d = _norm(data)
    sets = ", ".join(f"{k}=?" for k in MEMBER_FIELDS)
    conn.execute(f"UPDATE members SET {sets}, updated_at=CURRENT_TIMESTAMP WHERE id=?",
                 [*[d[k] for k in MEMBER_FIELDS], member_id])
    conn.commit()


def delete_member(conn, member_id: int):
    conn.execute("DELETE FROM payments WHERE member_id=?", (member_id,))
    conn.execute("DELETE FROM members WHERE id=?", (member_id,))
    conn.commit()


def members(conn) -> pd.DataFrame:
    """Members joined with their university details."""
    df = pd.read_sql(
        """SELECT m.*, u.name_en AS university, u.name_cn AS university_cn, u.ownership,
                  u.district, u.address AS university_address, u.postal_code
           FROM members m LEFT JOIN universities u ON u.id = m.university_id
           ORDER BY m.id""", conn)
    if df.empty:
        df = pd.DataFrame(columns=MEMBER_COLS + ["university", "university_cn", "ownership", "district",
                                                 "university_address", "postal_code"])
    return df


def get_member(conn, member_id: int) -> dict | None:
    df = members(conn)
    row = df[df["id"] == member_id]
    return None if row.empty else row.iloc[0].to_dict()


def approve_member(conn, member_id: int, join_date=None):
    join_date = (join_date or dt.date.today())
    if isinstance(join_date, (dt.date, dt.datetime)):
        join_date = join_date.isoformat()[:10]
    conn.execute("UPDATE members SET status='Active', join_date=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
                 (join_date, member_id))
    conn.commit()


def portal_lookup(conn, member_code: str, phone_last4: str) -> dict | None:
    """Self-service lookup: member code + last 4 digits of the registered phone number."""
    code = (member_code or "").strip().upper()
    last4 = "".join(ch for ch in (phone_last4 or "") if ch.isdigit())
    if not code or len(last4) != 4:
        return None
    df = members(conn)
    hit = df[df["member_code"].str.upper() == code]
    if hit.empty:
        return None
    m = hit.iloc[0].to_dict()
    digits = "".join(ch for ch in str(m.get("phone") or "") if ch.isdigit())
    return m if len(digits) >= 4 and digits[-4:] == last4 else None


# --------------------------------------------------------------------------------------
# QR code
# --------------------------------------------------------------------------------------
def member_link(base_url: str) -> str:
    base = (base_url or "").strip().rstrip("/")
    if base and not base.startswith(("http://", "https://")):
        base = "https://" + base
    return f"{base}/?page=member" if base else ""


def qr_image(data: str, color: str = "#0D2B55"):
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=14, border=3)
    qr.add_data(data)
    qr.make(fit=True)
    return qr.make_image(fill_color=color, back_color="white").convert("RGB")


def _font(size: int, bold=False):
    from PIL import ImageFont
    names = (["DejaVuSans-Bold.ttf", "Arial Bold.ttf", "arialbd.ttf"] if bold
             else ["DejaVuSans.ttf", "Arial.ttf", "arial.ttf"])
    for n in names:
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def qr_poster(qr_img, org_name: str, url: str, headline="Scan to join or check your dues") -> bytes:
    """A4-proportioned printable poster (PNG) with the QR code in the middle."""
    from PIL import Image, ImageDraw
    W, H = 1240, 1754
    navy, gold = (13, 43, 85), (201, 168, 76)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 300], fill=navy)
    d.rectangle([0, 300, W, 316], fill=gold)

    def center(text, y, font, fill):
        w = d.textlength(text, font=font)
        d.text(((W - w) / 2, y), text, font=font, fill=fill)

    def fit(text, size, bold, maxw):
        f = _font(size, bold)
        while d.textlength(text, font=f) > maxw and size > 18:
            size -= 2
            f = _font(size, bold)
        return f

    center(org_name, 95, fit(org_name, 64, True, W - 120), "white")
    center("Member Portal", 190, _font(40), gold)
    center(headline, 400, fit(headline, 56, True, W - 120), navy)
    q = qr_img.resize((820, 820), Image.NEAREST)
    x, y = (W - 820) // 2, 520
    d.rounded_rectangle([x - 30, y - 30, x + 850, y + 850], radius=30, outline=gold, width=6)
    img.paste(q, (x, y))
    center("1. Open your phone camera or WeChat scanner", 1430, _font(36), (55, 65, 81))
    center("2. Point it at the code and tap the link", 1485, _font(36), (55, 65, 81))
    center(url, 1600, fit(url, 30, False, W - 120), (107, 114, 128))
    buf = io.BytesIO()
    img.save(buf, format="PNG", dpi=(150, 150))
    return buf.getvalue()


def png_bytes(img) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# --------------------------------------------------------------------------------------
# Payments
# --------------------------------------------------------------------------------------
def add_payments(conn, member_id: int, periods: list[str], amount_each: float, payment_date,
                 method=None, reference=None, received_by=None, notes=None) -> int:
    if isinstance(payment_date, (dt.date, dt.datetime)):
        payment_date = payment_date.isoformat()[:10]
    for p in periods:
        conn.execute("""INSERT INTO payments(member_id, period, amount, payment_date, method, reference,
                        received_by, notes) VALUES (?,?,?,?,?,?,?,?)""",
                     (member_id, p, float(amount_each), payment_date, method, reference, received_by, notes))
    conn.commit()
    return len(periods)


def delete_payment(conn, payment_id: int):
    conn.execute("DELETE FROM payments WHERE id=?", (payment_id,))
    conn.commit()


def payments(conn) -> pd.DataFrame:
    df = pd.read_sql(
        """SELECT p.*, m.member_code, m.full_name, u.name_en AS university
           FROM payments p JOIN members m ON m.id = p.member_id
           LEFT JOIN universities u ON u.id = m.university_id
           ORDER BY p.payment_date DESC, p.id DESC""", conn)
    if df.empty:
        df = pd.DataFrame(columns=PAYMENT_COLS + ["member_code", "full_name", "university"])
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce").fillna(0.0)
    return df


# --------------------------------------------------------------------------------------
# Dues engine
# --------------------------------------------------------------------------------------
def period_of(d) -> str:
    if isinstance(d, str):
        return d[:7]
    return f"{d.year:04d}-{d.month:02d}"


def period_range(start: str, end: str) -> list[str]:
    """Inclusive list of 'YYYY-MM' from start to end."""
    if not start or not end or start > end:
        return []
    y, m = int(start[:4]), int(start[5:7])
    ey, em = int(end[:4]), int(end[5:7])
    out = []
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def shift_period(p: str, n: int) -> str:
    y, m = int(p[:4]), int(p[5:7])
    idx = y * 12 + (m - 1) + n
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def period_label(p: str) -> str:
    return dt.date(int(p[:4]), int(p[5:7]), 1).strftime("%b %Y")


def billable_periods(member: dict, upto: str) -> list[str]:
    """Months a member owes dues for, from joining month to `upto` (or exit month if earlier)."""
    if member.get("status") == PENDING:  # not billed until an admin approves them
        return []
    start = period_of(member["join_date"]) if member.get("join_date") else None
    end = upto
    exit_d = member.get("exit_date")
    if exit_d and isinstance(exit_d, str) and exit_d.strip():
        end = min(end, period_of(exit_d))
    return period_range(start, end) if start else []


def member_ledger(member: dict, pay_df: pd.DataFrame, upto: str) -> pd.DataFrame:
    """Month-by-month statement for one member."""
    due = float(member.get("monthly_due") or 0)
    periods = billable_periods(member, upto)
    mp = pay_df[pay_df["member_id"] == member["id"]]
    paid_by = mp.groupby("period")["amount"].sum().to_dict()
    # include prepaid / extra months that fall outside the billable window
    extra = sorted(p for p in paid_by if p not in periods)
    rows, running = [], 0.0
    for p in sorted(set(periods) | set(extra)):
        d = due if p in periods else 0.0
        paid = float(paid_by.get(p, 0.0))
        bal = d - paid
        running += bal
        if d == 0 and paid > 0:
            status = "Advance / Extra"
        elif d == 0:
            status = "No due"
        elif paid >= d:
            status = "Paid"
        elif paid > 0:
            status = "Partial"
        else:
            status = "Unpaid"
        rows.append({"Period": p, "Month": period_label(p), "Due": d, "Paid": paid,
                     "Balance": bal, "Running Balance": running, "Status": status})
    return pd.DataFrame(rows, columns=["Period", "Month", "Due", "Paid", "Balance", "Running Balance", "Status"])


def dues_summary(conn, upto: str, mem_df: pd.DataFrame | None = None,
                 pay_df: pd.DataFrame | None = None) -> pd.DataFrame:
    """One row per member: expected vs paid up to `upto`, balance, months in arrears."""
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    cols = ["id", "member_code", "full_name", "university", "district", "phone", "email",
            "membership_type", "status", "monthly_due", "months_billed", "expected", "paid",
            "balance", "months_in_arrears", "oldest_unpaid", "last_payment"]
    rows = []
    for m in mem_df.to_dict("records"):
        led = member_ledger(m, pay_df, upto)
        led_upto = led[led["Period"] <= upto]
        mp = pay_df[pay_df["member_id"] == m["id"]]
        expected = float(led_upto["Due"].sum())
        paid = float(mp["amount"].sum())  # all payments count as credit, incl. advance
        unpaid = led_upto[(led_upto["Due"] > 0) & (led_upto["Paid"] < led_upto["Due"])]
        rows.append({
            "id": m["id"], "member_code": m["member_code"], "full_name": m["full_name"],
            "university": m.get("university"), "district": m.get("district"),
            "phone": m.get("phone"), "email": m.get("email"),
            "membership_type": m.get("membership_type"), "status": m.get("status"),
            "monthly_due": float(m.get("monthly_due") or 0),
            "months_billed": int((led_upto["Due"] > 0).sum()),
            "expected": expected, "paid": paid, "balance": expected - paid,
            "months_in_arrears": len(unpaid),
            "oldest_unpaid": period_label(unpaid["Period"].min()) if len(unpaid) else "",
            "last_payment": mp["payment_date"].max() if len(mp) else "",
        })
    return pd.DataFrame(rows, columns=cols)


def monthly_collection(conn, period: str, mem_df=None, pay_df=None) -> pd.DataFrame:
    """For a single month: who owed what, who paid, who didn't."""
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    pp = pay_df[pay_df["period"] == period].groupby("member_id")["amount"].sum().to_dict()
    rows = []
    for m in mem_df.to_dict("records"):
        due = float(m.get("monthly_due") or 0) if period in billable_periods(m, period) else 0.0
        paid = float(pp.get(m["id"], 0.0))
        if due == 0 and paid == 0:
            continue
        status = "Paid" if paid >= due and due > 0 else "Partial" if paid > 0 and paid < due else \
            "Advance / Extra" if due == 0 else "Unpaid"
        rows.append({"member_code": m["member_code"], "full_name": m["full_name"],
                     "university": m.get("university"), "district": m.get("district"),
                     "phone": m.get("phone"), "membership_type": m.get("membership_type"),
                     "due": due, "paid": paid, "balance": due - paid, "status": status})
    return pd.DataFrame(rows, columns=["member_code", "full_name", "university", "district", "phone",
                                       "membership_type", "due", "paid", "balance", "status"])


def collection_trend(conn, end: str, months: int = 12, mem_df=None, pay_df=None) -> pd.DataFrame:
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    rows = []
    for p in period_range(shift_period(end, -(months - 1)), end):
        exp = sum(float(m.get("monthly_due") or 0) for m in mem_df.to_dict("records")
                  if p in billable_periods(m, p))
        col = float(pay_df.loc[pay_df["period"] == p, "amount"].sum())
        rows.append({"Period": p, "Month": period_label(p), "Expected": exp, "Collected": col,
                     "Outstanding": max(exp - col, 0)})
    return pd.DataFrame(rows)


def unpaid_periods(member: dict, pay_df: pd.DataFrame, upto: str) -> list[str]:
    led = member_ledger(member, pay_df, upto)
    return led[(led["Due"] > 0) & (led["Paid"] < led["Due"])]["Period"].tolist()


# --------------------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------------------
def to_excel(sheets: dict[str, pd.DataFrame], title: str | None = None) -> bytes:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    head_fill = PatternFill("solid", fgColor="0D2B55")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="0D2B55")
    thin = Side(style="thin", color="D0D5DD")
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            start = 2 if title else 0
            df.to_excel(xw, sheet_name=name[:31], index=False, startrow=start)
            ws = xw.sheets[name[:31]]
            if title:
                ws.cell(row=1, column=1, value=f"{title} — {name}").font = title_font
            hr = start + 1
            for ci, col in enumerate(df.columns, 1):
                c = ws.cell(row=hr, column=ci)
                c.fill, c.font = head_fill, head_font
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500).tolist()]) + 2
                ws.column_dimensions[get_column_letter(ci)].width = min(max(width, 10), 50)
                is_money = pd.api.types.is_float_dtype(df[col])
                for r in range(hr + 1, hr + 1 + len(df)):
                    cell = ws.cell(row=r, column=ci)
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    if is_money:
                        cell.number_format = "#,##0.00"
            ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return buf.getvalue()
